#!/usr/bin/env python3
"""
Validate that data/items.json matches the survey items in the source docx file.

Independently extracts items from the docx using the same rules:
- Heading 2 paragraphs within Appendix sections identify scales
- List Paragraph items under each Heading 2 are the survey items
- Trailing bracketed notes like [was "..."] are stripped
"""

import json
import re
from pathlib import Path
from docx import Document


def extract_items_from_docx(docx_path):
    """Extract 43 items from docx file. Returns list of (scale, text) tuples in order."""

    doc = Document(docx_path)

    # Build mapping from Heading 2 texts to scale IDs by scanning for Appendix sections
    scale_mapping = {}
    current_appendix = None

    for para in doc.paragraphs:
        if para.style.name == "Heading 1":
            text = para.text.strip()
            if text.startswith("Appendix A"):
                current_appendix = "A"
            elif text.startswith("Appendix B"):
                current_appendix = "B"
            elif text.startswith("Appendix C"):
                current_appendix = "C"

        elif para.style.name == "Heading 2" and current_appendix:
            text = para.text.strip()
            # Determine scale ID from heading text by prefix matching
            if text.startswith("A1"):
                scale_id = "A1"
            elif text.startswith("A2"):
                scale_id = "A2"
            elif text.startswith("A3"):
                scale_id = "A3"
            elif text.startswith("A4"):
                scale_id = "A4"
            elif text.startswith("A5"):
                scale_id = "A5"
            elif text.startswith("A6"):
                scale_id = "A6"
            elif text.startswith("D1"):
                scale_id = "D1"
            elif text.startswith("C1"):
                scale_id = "C1"
            elif text.startswith("C2"):
                scale_id = "C2"
            elif text.startswith("Co-creation"):
                scale_id = "CC"
            elif text.startswith("B1"):
                scale_id = "B1"
            elif text.startswith("B2"):
                scale_id = "B2"
            elif text.startswith("B3"):
                scale_id = "B3"
            else:
                scale_id = None

            if scale_id:
                scale_mapping[text] = scale_id

    # Extract items
    items = []
    current_scale = None

    for para in doc.paragraphs:
        if para.style.name == "Heading 2":
            text = para.text.strip()
            if text in scale_mapping:
                current_scale = scale_mapping[text]

        elif para.style.name == "List Paragraph" and current_scale:
            # Extract item text
            text = para.text.strip()

            # Strip trailing bracketed notes (both [was "..."] and generic [...])
            text = re.sub(r'\s*\[was\s+[^\]]*\]\s*$', '', text)
            text = re.sub(r'\s*\[[^\]]*\]\s*$', '', text)
            text = text.strip()

            items.append((current_scale, text))

    return items


def main():
    # Resolve paths relative to repo root
    repo_root = Path(__file__).resolve().parents[1]
    docx_path = repo_root / "source" / "Draft_Innovation_Capability_Framework-v02.docx"
    items_json_path = repo_root / "data" / "items.json"

    # Extract from docx
    docx_items = extract_items_from_docx(docx_path)

    # Load from JSON
    with open(items_json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    json_items = data['items']

    # Validate
    success = True
    errors = []

    if len(docx_items) != len(json_items):
        errors.append(f"Item count mismatch: docx has {len(docx_items)}, JSON has {len(json_items)}")
        success = False

    if len(docx_items) != 43:
        errors.append(f"Docx has {len(docx_items)} items, expected 43")
        success = False

    if len(json_items) != 43:
        errors.append(f"JSON has {len(json_items)} items, expected 43")
        success = False

    # Compare each item
    for i, (docx_item, json_item) in enumerate(zip(docx_items, json_items)):
        docx_scale, docx_text = docx_item
        json_scale = json_item['scale']
        json_text = json_item['text']

        if docx_scale != json_scale:
            errors.append(f"Index {i}: scale mismatch - docx has '{docx_scale}', JSON has '{json_scale}'")
            success = False

        if docx_text != json_text:
            errors.append(f"Index {i} ({json_item['id']}): text mismatch")
            errors.append(f"  Expected: {repr(docx_text)}")
            errors.append(f"  Got:      {repr(json_text)}")
            success = False

    # Print results
    if errors:
        for error in errors:
            print(error)
        return 1
    else:
        print("OK: 43 items match the document, in order.")
        return 0


if __name__ == "__main__":
    exit(main())
